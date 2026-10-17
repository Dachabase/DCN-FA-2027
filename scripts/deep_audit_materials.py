import os
import sys
import re
import json
import hashlib
import time

try:
    import pymupdf as fitz
except ImportError:
    try:
        import fitz
    except ImportError:
        fitz = None

def to_win_long_path(p):
    abs_p = os.path.abspath(p)
    if os.name == 'nt' and not abs_p.startswith('\\\\?\\'):
        return '\\\\?\\' + abs_p
    return abs_p

try:
    import docx
except ImportError:
    docx = None

try:
    import pptx
except ImportError:
    pptx = None

FACULTY_DIRS = [
    "Faculty of Art",
    "Faculty of Law",
    "Faculty of computing",
    "Faculty of economics",
    "Faculty of renewable natural resources",
    "Faculty of science and technology",
    "Faculty of social sciences",
    "Faculty of vet medicine",
]

COURSE_CODE_PATTERN = re.compile(r'\b([A-Z]{2,4}|C-[A-Z]{2,4})\s*([1-9][0-9]{2}[A-Z]?)\b', re.IGNORECASE)
YEAR_PATTERN = re.compile(r'\b(19\d\d|20\d\d)(?:/(\d{2,4}))?\b')

def get_file_hash(filepath, block_size=65536):
    hasher = hashlib.sha256()
    try:
        with open(to_win_long_path(filepath), 'rb') as f:
            for block in iter(lambda: f.read(block_size), b''):
                hasher.update(block)
        return hasher.hexdigest()
    except Exception:
        return None

def normalize_course_code(code_str):
    if not code_str:
        return None
    code_str = code_str.upper().strip()
    # Replace multiple spaces
    code_str = re.sub(r'\s+', ' ', code_str)
    # Match standard code
    m = re.match(r'^(C-)?([A-Z]{2,4})\s*(\d{3}[A-Z]?)$', code_str)
    if m:
        prefix = m.group(1) or ""
        subj = m.group(2)
        num = m.group(3)
        return f"{prefix}{subj} {num}"
    return code_str

def extract_pdf_info(filepath):
    if not fitz:
        return {"error": "PyMuPDF not installed"}
    try:
        doc = fitz.open(to_win_long_path(filepath))
        page_count = len(doc)
        total_text_len = 0
        pages_with_text = 0
        extracted_text_sample = []
        full_text_list = []
        
        for i in range(page_count):
            page = doc[i]
            t = page.get_text()
            full_text_list.append(t)
            t_clean = t.strip()
            total_text_len += len(t_clean)
            if len(t_clean) > 30:
                pages_with_text += 1
            if i < 3 and t_clean:
                extracted_text_sample.append(t_clean[:500])
                
        is_scanned = (pages_with_text == 0 and page_count > 0) or (total_text_len / max(1, page_count) < 20)
        
        full_text = "\n".join(full_text_list)
        doc.close()
        return {
            "page_count": page_count,
            "is_scanned": is_scanned,
            "text_char_count": total_text_len,
            "pages_with_text": pages_with_text,
            "text_sample": "\n---\n".join(extracted_text_sample)[:1000],
            "full_text_snippet": full_text[:2000],
            "full_text": full_text
        }
    except Exception as e:
        return {"error": str(e), "page_count": 0, "is_scanned": False}

def extract_docx_info(filepath):
    try:
        if docx:
            doc = docx.Document(to_win_long_path(filepath))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            full_text = "\n".join(paragraphs)
            return {
                "page_count": len(paragraphs) // 10 + 1,
                "is_scanned": False,
                "text_char_count": len(full_text),
                "text_sample": full_text[:1000],
                "full_text": full_text
            }
        else:
            return {"error": "docx library not available"}
    except Exception as e:
        return {"error": str(e), "is_scanned": False}

def extract_pptx_info(filepath):
    try:
        if pptx:
            prs = pptx.Presentation(to_win_long_path(filepath))
            slide_texts = []
            for slide in prs.slides:
                slide_text = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_text.append(shape.text.strip())
                if slide_text:
                    slide_texts.append("\n".join(slide_text))
            full_text = "\n\n--- SLIDE ---\n\n".join(slide_texts)
            return {
                "page_count": len(prs.slides),
                "is_scanned": False,
                "text_char_count": len(full_text),
                "text_sample": full_text[:1000],
                "full_text": full_text
            }
        else:
            return {"error": "pptx library not available"}
    except Exception as e:
        return {"error": str(e), "is_scanned": False}

def detect_category(path_lower, filename_lower, text_lower=""):
    combined = f"{path_lower} {filename_lower} {text_lower[:1000]}"
    if any(k in combined for k in ["past question", "past-question", "past_question", "pq", "examination", "exam question", "test_", "test 1", "test 2"]):
        return "PAST_QUESTIONS"
    if any(k in combined for k in ["slide", "lecture note", "handout", "presentation", "module", "chapter", "lecture"]):
        return "LECTURE_NOTES"
    if any(k in combined for k in ["textbook", "book", "author", "edition"]):
        return "TEXTBOOK"
    if any(k in combined for k in ["practical", "manual", "lab", "laboratory", "experiment"]):
        return "PRACTICAL_MANUAL"
    if any(k in combined for k in ["tutorial", "assignment", "problem set", "exercise", "homework"]):
        return "TUTORIAL_SHEET"
    if any(k in combined for k in ["syllabus", "course outline", "curriculum"]):
        return "SYLLABUS"
    return "ACADEMIC_MATERIAL"

def detect_course(rel_path, filename, text=""):
    # 1. Search in filename and directory path
    path_matches = COURSE_CODE_PATTERN.findall(rel_path)
    if path_matches:
        # e.g. [('COS', '101'), ...]
        for m in path_matches:
            code = normalize_course_code(f"{m[0]} {m[1]}")
            num = int(re.sub(r'\D', '', m[1]))
            if 100 <= num <= 199: # 100 level
                return code
            
    # 2. Search in text content (first 2000 chars)
    if text:
        text_matches = COURSE_CODE_PATTERN.findall(text[:2000])
        for m in text_matches:
            code = normalize_course_code(f"{m[0]} {m[1]}")
            num = int(re.sub(r'\D', '', m[1]))
            if 100 <= num <= 199:
                return code
                
    return None

def detect_semester(rel_path, filename, course_code="", text=""):
    comb = f"{rel_path} {filename} {text[:1000]}".lower()
    if any(k in comb for k in ["first semester", "1st semester", "first sem", "1st sem", "harmattan"]):
        return "1st"
    if any(k in comb for k in ["second semester", "2nd semester", "second sem", "2nd sem", "rain"]):
        return "2nd"
    if course_code:
        # Check course number last digit: odd = 1st, even = 2nd
        m = re.search(r'\d{3}', course_code)
        if m:
            last_digit = int(m.group(0)[-1])
            return "1st" if last_digit % 2 == 1 else "2nd"
    return "1st"

def detect_exam_year(rel_path, filename, text=""):
    comb = f"{filename} {rel_path} {text[:2000]}"
    matches = YEAR_PATTERN.findall(comb)
    for m in matches:
        year = int(m[0])
        if 1995 <= year <= 2026:
            return year
    return None

def run_deep_audit():
    start_time = time.time()
    all_files = []
    
    print("Beginning Deep Academic Materials Audit across all 8 faculties...")
    
    for fac_dir in FACULTY_DIRS:
        if not os.path.exists(fac_dir):
            print(f"Directory {fac_dir} not found! Skipping.")
            continue
            
        print(f"\nScanning: {fac_dir}")
        count_in_fac = 0
        
        base_root = to_win_long_path('.')
        for dirpath, _, filenames in os.walk(to_win_long_path(fac_dir)):
            for fname in filenames:
                full_path = os.path.join(dirpath, fname)
                try:
                    rel_path = os.path.relpath(full_path, base_root)
                except Exception:
                    rel_path = full_path.replace(base_root, '').lstrip('\\/')
                ext = os.path.splitext(fname)[1].lower()
                try:
                    size = os.path.getsize(to_win_long_path(full_path))
                except Exception:
                    size = 0
                
                count_in_fac += 1
                all_files.append({
                    "faculty_dir": fac_dir,
                    "rel_path": rel_path,
                    "filename": fname,
                    "ext": ext,
                    "size": size,
                    "full_path": full_path
                })
        print(f"Found {count_in_fac} files in {fac_dir}")
        
    print(f"\nTotal files discovered across all faculties: {len(all_files)}")
    
    # Process each file
    processed_inventory = []
    course_counts = {}
    category_counts = {}
    format_counts = {}
    scanned_count = 0
    past_questions_count = 0
    
    for idx, item in enumerate(all_files):
        if (idx + 1) % 100 == 0 or idx == len(all_files) - 1:
            print(f"Processed {idx + 1}/{len(all_files)} files ({(idx+1)/len(all_files)*100:.1f}%)...")
            
        ext = item["ext"]
        format_counts[ext] = format_counts.get(ext, 0) + 1
        
        file_hash = get_file_hash(item["full_path"])
        
        extraction = {}
        if ext == ".pdf":
            extraction = extract_pdf_info(item["full_path"])
        elif ext in [".docx", ".doc"]:
            extraction = extract_docx_info(item["full_path"])
        elif ext in [".pptx", ".ppt"]:
            extraction = extract_pptx_info(item["full_path"])
        elif ext in [".txt", ".md"]:
            try:
                with open(to_win_long_path(item["full_path"]), "r", encoding="utf-8", errors="ignore") as f:
                    txt = f.read()
                    extraction = {
                        "page_count": 1,
                        "is_scanned": False,
                        "text_char_count": len(txt),
                        "text_sample": txt[:1000],
                        "full_text": txt
                    }
            except Exception as e:
                extraction = {"error": str(e)}
        else:
            extraction = {
                "page_count": 1,
                "is_scanned": False,
                "text_char_count": 0,
                "text_sample": "",
                "full_text": ""
            }
            
        is_scanned = extraction.get("is_scanned", False)
        if is_scanned:
            scanned_count += 1
            
        full_text = extraction.get("full_text", "")
        # Remove full_text from stored item to keep json size manageable, but keep snippet
        full_text_snippet = full_text[:4000] if full_text else ""
        
        course_code = detect_course(item["rel_path"], item["filename"], full_text)
        category = detect_category(item["rel_path"].lower(), item["filename"].lower(), full_text.lower())
        semester = detect_semester(item["rel_path"], item["filename"], course_code, full_text)
        exam_year = detect_exam_year(item["rel_path"], item["filename"], full_text)
        
        if category == "PAST_QUESTIONS":
            past_questions_count += 1
            
        if course_code:
            course_counts[course_code] = course_counts.get(course_code, 0) + 1
        category_counts[category] = category_counts.get(category, 0) + 1
        
        # Check if contains questions
        has_questions = False
        text_lower = full_text.lower()
        if any(w in text_lower for w in ["question 1", "question 2", "q1.", "q2.", "1.", "attempt any", "answer all"]) and ("a)" in text_lower or "a." in text_lower or "option" in text_lower or "marks" in text_lower):
            has_questions = True
            
        record = {
            "id": f"mat_{idx+1:05d}",
            "filename": item["filename"],
            "rel_path": item["rel_path"],
            "faculty_dir": item["faculty_dir"],
            "ext": ext,
            "size": item["size"],
            "hash": file_hash,
            "page_count": extraction.get("page_count", 0),
            "is_scanned": is_scanned,
            "char_count": extraction.get("text_char_count", 0),
            "course_code": course_code,
            "category": category,
            "semester": semester,
            "exam_year": exam_year,
            "has_questions": has_questions,
            "text_snippet": full_text_snippet,
            "status": "PROCESSED" if not extraction.get("error") else "EXTRACTION_ERROR",
            "error": extraction.get("error")
        }
        processed_inventory.append(record)
        
    duration = time.time() - start_time
    print(f"\nAudit completed in {duration:.1f} seconds!")
    print(f"Total files audited: {len(processed_inventory)}")
    print(f"Scanned/Image-only documents: {scanned_count}")
    print(f"Documents identified as Past Questions: {past_questions_count}")
    print(f"Distinct course codes mapped: {len(course_counts)}")
    
    print("\nTop 25 Course Codes Found:")
    for c, cnt in sorted(course_counts.items(), key=lambda x: -x[1])[:25]:
        print(f"  {c}: {cnt} files")
        
    print("\nCategories:")
    for cat, cnt in sorted(category_counts.items(), key=lambda x: -x[1]):
        print(f"  {cat}: {cnt} files")
        
    print("\nFormats:")
    for fmt, cnt in sorted(format_counts.items(), key=lambda x: -x[1]):
        print(f"  {fmt}: {cnt} files")
        
    os.makedirs("data", exist_ok=True)
    out_path = "data/master_materials_inventory.json"
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(processed_inventory, f, indent=2, ensure_ascii=False)
    print(f"\nSaved master inventory to {out_path} ({os.path.getsize(out_path)/1024/1024:.2f} MB)")

if __name__ == "__main__":
    run_deep_audit()
